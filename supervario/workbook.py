"""Human-readable Leapfrog estimator recommendations."""

import re
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from .recommendations import pass_guidance


def recommendation_workbook(results):
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    headers = [
        "Domain", "Recommendation status", "Orientation status", "Kriging type", "Variogram Model",
        "Pass", "Maximum", "Intermed.", "Minimum", "Dip", "Dip Azimuth", "Pitch",
        "Minimum samples", "Maximum samples", "Sample-probe support", "Apply to", "Adoption",
        "Mean value", "Discretisation X/Y/Z", "Variable Orientation", "Clip input values",
        "Outlier Restriction", "Sector Search", "Drillhole Limit", "Warnings",
    ]
    summary.append(headers)
    used = {"summary"}

    def append(sheet, values):
        sheet.append(values)
        # Domain names and notes are text, even when they start with '='.
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"

    def flatten(sheet, section, values, prefix=""):
        for key, value in values.items():
            field = f"{prefix} / {key}" if prefix else key
            if isinstance(value, dict):
                flatten(sheet, section, value, field)
            else:
                if isinstance(value, list):
                    value = ", ".join(map(str, value))
                if value is None:
                    value = "Not specified / requires review"
                append(sheet, [section, field, value])

    for result in results:
        advice = result["estimator_recommendations"]
        parameters = advice["leapfrog_parameters"]
        domain = str(result["domain"])
        stem = re.sub(r"[\x00-\x1f\\/*?:\[\]]", "_", domain).strip("'")[:31] or "Domain"
        name, number = stem, 2
        while name.casefold() in used:
            suffix = f"_{number}"
            name = stem[:31 - len(suffix)] + suffix
            number += 1
        used.add(name.casefold())
        sheet = workbook.create_sheet(name)
        append(sheet, ["Section", "Leapfrog field / item", "Recommendation / explanation"])
        append(sheet, ["Overview", "Domain", domain])
        append(sheet, ["Overview", "Recommendation status", advice["status"].replace("_", " ")])
        append(sheet, ["Overview", "Orientation status", advice["orientation_status"].replace("_", " ")])
        append(sheet, ["Overview", "Scope", parameters["status"]])
        append(sheet, ["Overview", "Distances", "Input coordinate units; centre-to-boundary ranges, NOT diameters."])
        append(sheet, ["Overview", "Pass count", advice.get("recommended_pass_count", 0)])
        for index, item in enumerate(pass_guidance(advice), 1):
            append(sheet, ["Pass guidance", str(index), item])
        for section in ("Interpolant", "Value Clipping", "Search", "Outputs"):
            flatten(sheet, section, parameters[section])
        for entry in parameters["passes"]:
            flatten(sheet, f"Pass {entry['pass']}", entry)
        if not parameters["passes"]:
            append(sheet, ["Passes", "Unavailable", "No supported search ellipsoid dimensions; recalculate/review evidence."])
        for section, items in (("Warnings", advice["warnings"]), ("Validation", advice["validation"])):
            for index, item in enumerate(items, 1):
                append(sheet, [section, str(index), item])
        flatten(sheet, "Comparisons", advice["comparisons"])
        append(sheet, ["Reference", "Leapfrog help", parameters["source"]])
        for p, entry in zip(advice["passes"], parameters["passes"]):
            ellipsoid = entry["Ellipsoid"]
            append(summary, [
                domain, advice["status"].replace("_", " "), advice["orientation_status"].replace("_", " "),
                parameters["Interpolant"]["Kriging type"], advice["variogram_name"], p["pass"],
                *[ellipsoid[k] for k in ("Maximum", "Intermed.", "Minimum", "Dip", "Dip Azimuth", "Pitch")],
                p["minimum_samples"], p["maximum_samples"], p["sample_probe_support_fraction"],
                p["apply_to"], p["adoption_status"], "Not applicable (Ordinary)", "Requires block validation",
                "Off", "Off; bounds require review", "Off; thresholds require review",
                "Off; test Octant 3 samples/sector, 4 empty sectors",
                "Off pending drillhole IDs / validation", "\n".join(advice["warnings"]),
            ])
        if not advice["passes"]:
            append(summary, [domain, advice["status"].replace("_", " "),
                             advice["orientation_status"].replace("_", " "),
                             parameters["Interpolant"]["Kriging type"], advice["variogram_name"],
                             None, *[None] * 18, "\n".join(advice["warnings"])])
        sheet.freeze_panes = "C2"
        sheet.column_dimensions["A"].width = 22
        sheet.column_dimensions["B"].width = 48
        sheet.column_dimensions["C"].width = 100

    for sheet in workbook:
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="234E70")
        sheet.row_dimensions[1].height = 32
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="EDF3F8")
                if isinstance(cell.value, float):
                    cell.number_format = "0.######"
            longest = max((len(str(c.value or "")) for c in row), default=0)
            sheet.row_dimensions[row[0].row].height = min(150, max(30, 15 * (1 + longest // 90)))
        sheet.print_title_rows = "1:1"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
    summary.freeze_panes = "F2"
    for column in summary.columns:
        summary.column_dimensions[column[0].column_letter].width = 22
    for column in ("B", "C", "E", "P", "Q", "W", "X", "Y"):
        summary.column_dimensions[column].width = 45
    for cell in summary["O"][1:]:
        cell.number_format = "0.0%"
    if summary.max_row > 1:
        table = Table(displayName="EstimatorSummary", ref=summary.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        summary.add_table(table)
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()
