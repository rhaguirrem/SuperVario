from io import BytesIO

import numpy as np
from openpyxl import load_workbook

from supervario.recommendations import evidence, recommend
from supervario.workbook import recommendation_workbook
from test_recommendations import result_with_ranges


def result(domain, supported=True):
    item = result_with_ranges((20, 10, 3))
    item["domain"] = domain
    item["model"]["metadata"]["name"] = "Variogram " + domain
    if supported:
        item["neighbourhood_evidence"] = evidence(np.random.default_rng(2).random((30, 3)),
                                                  np.zeros((1, 3)))
    item["estimator_recommendations"] = recommend(item)
    return item


def test_workbook_matches_pass_values_and_has_readable_domain_details():
    results = [result("UGB01"), result("UGB02", False)]
    workbook = load_workbook(BytesIO(recommendation_workbook(results)))
    assert workbook.sheetnames == ["Summary", "UGB01", "UGB02"]
    summary = workbook["Summary"]
    headers = [c.value for c in summary[1]]
    rows = [dict(zip(headers, row)) for row in summary.iter_rows(min_row=2, values_only=True)]
    first = rows[0]
    p = results[0]["estimator_recommendations"]["passes"][0]
    assert first["Domain"] == "UGB01"
    assert first["Maximum"] == p["semi_axes"]["major"]
    assert first["Intermed."] == p["semi_axes"]["semi_major"]
    assert first["Minimum"] == p["semi_axes"]["minor"]
    assert first["Minimum samples"] == 8
    assert first["Maximum samples"] == 24
    assert first["Sample-probe support"] == p["sample_probe_support_fraction"]
    assert summary["O2"].number_format == "0.0%"
    assert rows[-1]["Domain"] == "UGB02"
    assert rows[-1]["Pass"] is None
    assert rows[-1]["Recommendation status"] == "support evidence unavailable"
    assert rows[-1]["Warnings"]
    assert summary.freeze_panes == "F2"
    assert "EstimatorSummary" in summary.tables
    details = list(workbook["UGB01"].values)
    assert any(row[:2] == ("Interpolant", "Kriging type") and row[2] == "Ordinary" for row in details)
    assert any(row[0] == "Warnings" for row in details)
    assert any(row[0] == "Validation" for row in details)
    assert [row[2] for row in details if row[0] == "Pass guidance"] == (
        results[0]["estimator_recommendations"]["pass_guidance"])
    assert any(row[1] == "Discretisation / X" and "Not specified" in row[2] for row in details)
    assert any(row[1] == "Ellipsoid / Maximum" and row[2] == first["Maximum"] for row in details)
    workbook.close()


def test_sheet_names_are_safe_unique_and_text_is_not_a_formula():
    domains = ["Summary", "summary", "a/b", "a:b", "x" * 50, "x" * 51, "=2+2", "''"]
    workbook = load_workbook(BytesIO(recommendation_workbook([result(d, False) for d in domains])))
    assert len(workbook.sheetnames) == len(domains) + 1
    assert len({name.casefold() for name in workbook.sheetnames}) == len(workbook.sheetnames)
    assert all(len(name) <= 31 for name in workbook.sheetnames)
    assert all(not any(c in name for c in r"\/:*?[]") for name in workbook.sheetnames)
    for domain, sheet in zip(domains, list(workbook)[1:]):
        assert sheet["C2"].value == domain
        assert sheet["C2"].data_type == "s"
    assert all(cell.data_type != "f" for sheet in workbook for row in sheet for cell in row)
    workbook.close()
