import copy
import json

import numpy as np
import pandas as pd
import pytest

from supervario import lfv
from supervario.cli import main
from supervario.data import DOMAIN
from supervario.pipeline import contacts, export_report, input_paths, load_config, run
from supervario.trends import string_theory_candidates


def test_drilling_metadata_flows_from_csv_to_domain_lag_diagnostics(project, tmp_path):
    path = project["samples"]["path"]
    frame = pd.read_csv(path)
    frame["HOLE"] = [f"{i % 20:03d}" for i in range(len(frame))]
    frame["FROM"] = np.arange(len(frame)) // 20
    frame["TO"] = frame["FROM"] + 1
    frame.to_csv(path, index=False)
    project["samples"].update(hole_id="HOLE", **{"from": "FROM", "to": "TO"})
    project["variogram"]["rotation"] = {"dip_direction": 0, "dip": 0, "pitch": 0}
    report = run(project, lambda _: None)
    assert not report["failures"]
    for result in report["results"]:
        info = result["lag_selection"]
        assert info["method"] == "spacing_informed_candidate_search"
        assert info["holes_checked"] == 20
        assert info["interval_length_p10_p50_p90"] == [1, 1, 1]
        assert len(result["radial"]["gamma"][0]) == info["bins"]
        assert result["downhole"]["status"] in ("available", "insufficient_support")
        assert len(result["downhole"]["hole_count"]) == 12
    export_report(report, tmp_path / "export")
    exported = json.loads((tmp_path / "export" / "analysis.json").read_text())
    assert exported["results"][0]["lag_selection"] == report["results"][0]["lag_selection"]
    assert exported["results"][0]["downhole"] == report["results"][0]["downhole"]

def test_readable_export_names_only_suffix_actual_collisions():
    from supervario.pipeline import result_filenames
    assert result_filenames(["UGB01", "UGB02A", "UGB02B"]) == [
        "variogram_UGB01.lfv", "variogram_UGB02A.lfv", "variogram_UGB02B.lfv"]
    domains = ["A/B", "A:B", "A_B_2", "ugB", "UGB", "."]
    names = result_filenames(domains)
    assert names[:3] == ["variogram_A_B.lfv", "variogram_A_B_3.lfv", "variogram_A_B_2.lfv"]
    assert len({name.casefold() for name in names}) == len(names)
    assert names[-1] == "variogram_domain.lfv"
    assert dict(zip(domains, names)) == dict(zip(reversed(domains), result_filenames(list(reversed(domains)))))
    assert len(set(result_filenames(["x" * 80, "x" * 81]))) == 2


def test_existing_excel_requires_approval_before_any_other_output_is_written(project, tmp_path):
    from test_workbook import result
    report = {"config": project, "application": "SuperVario", "version": "test",
              "results": [result("A")]}
    output = tmp_path / "output"
    output.mkdir()
    path = output / "estimator_recommendations.xlsx"
    path.write_bytes(b"existing workbook")
    with pytest.raises(FileExistsError, match="estimator_recommendations.xlsx"):
        export_report(report, output)
    assert list(output.iterdir()) == [path]
    assert path.read_bytes() == b"existing workbook"


@pytest.mark.parametrize("in_plane_step", [5, 15])
def test_full_pipeline_export_and_provenance(project, tmp_path, in_plane_step):
    project["variogram"]["in_plane_step"] = in_plane_step
    project["variogram"]["azimuth_step"] = 30 if in_plane_step == 5 else 15
    project["variogram"]["plunge_step"] = 45 if in_plane_step == 5 else 30
    report = run(project, lambda _: None)
    assert {r["domain"] for r in report["results"]} == {"A", "B"}
    assert report["failures"] == []
    assert all(r["borrowed_samples"] == 0 for r in report["results"])
    paths = export_report(report, tmp_path / "output")
    assert len(paths) == 5
    for path in paths[:-3]:
        model = lfv.load(path)
        assert model["metadata"]["attribute"] == "grade"
        assert model["sill"] == pytest.approx(model["nugget"] + sum(s["contribution"] for s in model["structures"]))
    with open(paths[-1]) as handle:
        saved = json.load(handle)
    assert saved["config"] == project
    assert saved["results"][0]["pairs"]["drawn"] <= 300000
    for result in saved["results"]:
        rotations = [s["anisotropy"]["rotation"] for s in result["model"]["structures"]]
        assert all(r == result["radial"]["rotation"] for r in rotations)
        assert result["orientation_diagnostics"]["status"] in (
            "major_direction_unresolved", "provisional_supported_direction")
        assert len(result["orientation_diagnostics"]["resampled_major_vectors"]) == 4
        assert result["orientation_diagnostics"]["azimuth_step_degrees"] == project["variogram"]["azimuth_step"]
        assert result["orientation_diagnostics"]["plunge_step_degrees"] == project["variogram"]["plunge_step"]
        refinement = result["fit"]["plane_refinement"]
        assert refinement["in_plane_step_degrees"] == in_plane_step
        assert sorted({r["angle_degrees"] for r in refinement["candidates"]
                       if r["source"] == "in_plane"}) == list(range(0, 90, in_plane_step))
        assert result["orientation_diagnostics"]["frame_shortlist"]["ordered_frames"] <= 24
        assert any(r["source"] == "alternative_3d" for r in refinement["candidates"])
        assert refinement == result["orientation_diagnostics"]["plane_refinement"]
        assert refinement["status"] in (
            "rotated_frame_selected", "alternative_3d_frame_selected", "baseline_retained", "baseline_validation_support_insufficient",
            "common_validation_distances_insufficient")
        for structure in result["model"]["structures"]:
            ranges = structure["anisotropy"]["ranges"]
            assert ranges["major"] >= ranges["semi_major"] >= ranges["minor"] > 0
            assert all(ranges[axis] <= result["fit"]["range_caps"][axis] for axis in ranges)
    with open(paths[-2]) as handle:
        advice = json.load(handle)
    from openpyxl import load_workbook
    workbook = load_workbook(paths[-3])
    assert workbook.sheetnames == ["Summary", "A", "B"]
    assert workbook["Summary"].max_row == 1 + sum(
        r["estimator_recommendations"]["recommended_pass_count"] for r in saved["results"])
    workbook.close()
    assert {d["domain"] for d in advice["domains"]} == {"A", "B"}
    assert all(1 <= d["recommended_pass_count"] <= 5 for d in advice["domains"])
    for domain in advice["domains"]:
        mapped = domain["leapfrog_parameters"]
        assert mapped["Interpolant"]["Kriging type"] == "Ordinary"
        assert len(mapped["passes"]) == domain["recommended_pass_count"]
        for p, entry in zip(domain["passes"], mapped["passes"]):
            assert entry["Ellipsoid"]["Maximum"] == p["semi_axes"]["major"]
            assert entry["Ellipsoid"]["Dip Azimuth"] == p["rotation"]["dip_direction"]
            assert entry["Search"]["Minimum samples"] == p["minimum_samples"]
        saved_result = next(r for r in saved["results"] if r["domain"] == domain["domain"])
        assert saved_result["estimator_recommendations"]["leapfrog_parameters"] == mapped
    with pytest.raises(FileExistsError):
        export_report(report, tmp_path / "output")
    export_report(report, tmp_path / "output", overwrite=True)


def test_manual_orientation_bypasses_automatic_plane_compromise(project, monkeypatch):
    def reject_refinement(*args, **kwargs):
        pytest.fail("Manual orientation must not be refined")

    monkeypatch.setattr("supervario.variography.refine_plane", reject_refinement)
    rotation = {"dip_direction": 105, "dip": 57, "pitch": 95}
    project["variogram"]["rotation"] = rotation
    report = run(project, lambda _: None)
    assert report["failures"] == []
    for result in report["results"]:
        assert result["model"]["structures"][0]["anisotropy"]["rotation"] == rotation
        assert result["orientation_diagnostics"]["status"] == "manual_direction_not_statistically_validated"
        assert "plane_refinement" not in result["fit"]


def test_soft_boundaries_are_opt_in_and_distance_limited(project):
    hard = run(project, lambda _: None)
    project["boundaries"]["approved_pairs"] = [["A", "B"]]
    soft = run(project, lambda _: None)
    assert all(r["borrowed_samples"] == 0 for r in hard["results"])
    assert all(0 < r["borrowed_samples"] < r["hard_domain_samples"] for r in soft["results"])
    for result in soft["results"]:
        assert result["sample_count"] == result["hard_domain_samples"] + result["borrowed_samples"]


def test_invalid_boundary_approval_is_rejected(project):
    project["boundaries"]["approved_pairs"] = [["A", "absent"]]
    with pytest.raises(ValueError, match="Invalid approved"):
        run(project, lambda _: None)


@pytest.mark.parametrize("key,value,label", [
    ("bins", 0, "Lag bins"), ("bins", 3.5, "whole number"),
    ("max_pairs", 0, "Pair cap"), ("max_samples", 0, "Sample cap"),
    ("min_pairs", 0, "Minimum pairs"), ("structures", 0, "Spherical structures"),
    ("angle_tolerance", 0, "Cone half-angle"), ("lag", -1, "Lag distance"),
    ("max_distance", float("nan"), "Maximum distance"),
])
def test_invalid_settings_fail_before_reading_inputs(project, key, value, label):
    project["variogram"][key] = value
    project["samples"]["path"] = "missing-samples.csv"
    with pytest.raises(ValueError, match=label):
        run(project, lambda _: None)


def test_contact_heuristics_do_not_autoshare_samples():
    frame = pd.DataFrame({"X": np.r_[np.arange(30), np.arange(30) + 0.1],
                          "Y": 0, "Z": 0, "grade": np.tile(np.arange(30), 2),
                          DOMAIN: ["A"] * 30 + ["B"] * 30})
    result = contacts(frame, ["X", "Y", "Z"], "grade", {"distance": 1, "min_samples": 20})
    assert result[0]["recommend_review"]
    assert result[0]["standardised_mean_difference"] == 0
    assert "proxy" in result[0]["method"]


def test_input_overwrite_protection_including_aliases(project, tmp_path):
    source = project["samples"]["path"]
    with pytest.raises(ValueError, match="overwrite"):
        lfv.write_json(source, {}, input_paths(project))
    alias = tmp_path / "alias.json"
    import os
    os.link(source, alias)
    with pytest.raises(ValueError, match="aliases"):
        lfv.write_json(alias, {}, [source])


def test_partial_domain_failure_is_not_a_fake_model(project):
    frame = pd.read_csv(project["samples"]["path"])
    frame.loc[frame["X"] >= 50, "grade"] = 3
    frame.to_csv(project["samples"]["path"], index=False)
    report = run(project, lambda _: None)
    assert [r["domain"] for r in report["results"]] == ["A"]
    assert report["failures"][0]["domain"] == "B"
    assert "zero" in report["failures"][0]["error"]


def test_relative_config_paths_are_resolved_to_config_folder(tmp_path):
    from supervario.pipeline import default_config
    config = default_config()
    config["samples"]["path"] = "samples.csv"
    config["domains"]["path"] = "blocks.csv"
    config["variogram"]["trend_file"] = "path_stats.csv"
    path = tmp_path / "project.json"
    lfv.write_json(path, config)
    loaded = load_config(path)
    assert loaded["samples"]["path"] == str(tmp_path / "samples.csv")
    assert loaded["variogram"]["trend_file"] == str(tmp_path / "path_stats.csv")


def test_string_theory_import_uses_accepted_matching_categories_only(tmp_path):
    path = tmp_path / "path_stats.csv"
    pd.DataFrame({
        "Azimuth": [90, 91, 270, 180], "Dip": [20, 21, 0, 30],
        "Category": ["A", "A", "A", "B"], "Accepted": [True, True, False, True],
    }).to_csv(path, index=False)
    candidates, audit = string_theory_candidates(str(path), "A")
    assert audit["accepted_paths"] == 2
    assert candidates[0] == {"azimuth": 95, "plunge": 25, "count": 2}
    with pytest.raises(ValueError, match="No accepted"):
        string_theory_candidates(str(path), "C")


@pytest.mark.parametrize("manual", [False, True])
def test_deprecated_string_theory_input_is_ignored_without_reading_and_reported(project, tmp_path, manual):
    if manual:
        project["variogram"]["rotation"] = {"dip_direction": 105, "dip": 57, "pitch": 95}
    baseline = run(project, lambda _: None)
    path = tmp_path / "missing_point_cloud_or_stats.csv"
    project["variogram"]["trend_file"] = str(path)
    logs = []
    report = run(project, logs.append)
    assert not path.exists()
    assert report["failures"] == []
    for result, original in zip(report["results"], baseline["results"]):
        assert result["model"] == original["model"]
        assert result["experiments"] == original["experiments"]
        assert result["string_theory"]["status"] == "deprecated_ignored"
        assert result["string_theory"]["source"] == str(path)
        assert any("deprecated and ignored" in w for w in result["fit"]["warnings"])
    assert any("deprecated and ignored" in message for message in logs)
    assert str(path) in input_paths(project)
    paths = export_report(report, tmp_path / "deprecated-output")
    with open(paths[-1]) as handle:
        saved = json.load(handle)
    assert all(r["string_theory"]["status"] == "deprecated_ignored" for r in saved["results"])


def test_export_preflights_invalid_report_without_partial_files(project, tmp_path):
    report = run(project, lambda _: None)
    report["config"]["bad"] = float("nan")
    output = tmp_path / "output"
    with pytest.raises(ValueError):
        export_report(report, output)
    assert not output.exists()


def test_export_refreshes_recommendations_after_model_changes(project, tmp_path):
    report = run(project, lambda _: None)
    result = report["results"][0]
    for structure in result["model"]["structures"]:
        structure["anisotropy"]["ranges"] = {"major": 4, "semi_major": 3, "minor": 2}
    paths = export_report(report, tmp_path / "edited")
    with open(paths[-2]) as handle:
        advice = json.load(handle)["domains"][0]
    assert advice["reference"]["base_semi_axes"] == {"major": 4, "semi_major": 3, "minor": 2}
    # Export operates on a snapshot and does not mutate the caller's report.
    assert result["estimator_recommendations"]["reference"]["base_semi_axes"] != advice["reference"]["base_semi_axes"]


def test_cli_header_and_reference_validation(project, capsys):
    from pathlib import Path
    assert main(["header", project["samples"]["path"]]) == 0
    assert "grade" in capsys.readouterr().out
    reference = Path(__file__).resolve().parents[1] / "Reference" / "Variogram Model RQDpc UGB01.lfv"
    assert main(["inspect-lfv", str(reference)]) == 0
    assert "normalised" in capsys.readouterr().out
